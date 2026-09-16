"""Import must read actual documents, preserve source and wait for publication."""

import asyncio
import io
import json
from urllib.parse import quote

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from sqlalchemy import select

from app.models import Job, Study, StudyVersion
from app.providers import TextResult, Usage
from app.worker import Worker
from test_core import STUDY, headers, new_study


def document_bytes(kind):
    stream = io.BytesIO()
    if kind == "docx":
        doc = Document()
        doc.add_paragraph("调研名称：材料流转体验")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "研究目标"
        table.cell(0, 1).text = "理解补充材料的具体经历"
        doc.add_paragraph("访谈主题：一次补件经历")
        doc.save(stream)
    elif kind == "xlsx":
        book = Workbook()
        book.active.title = "研究目标"
        book.active.append(["调研名称", "材料流转体验"])
        book.active.append(["研究目标", "理解补充材料的具体经历"])
        other = book.create_sheet("访谈提纲")
        other.append(["主题", "问题"])
        other.append(["一次补件经历", "当时如何处理？"])
        other["C3"] = "=1+1"
        book.save(stream)
    elif kind == "pptx":
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "材料流转体验"
        second = deck.slides.add_slide(deck.slide_layouts[6])
        table = second.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(5), Inches(2)).table
        table.cell(0, 0).text = "研究目标"
        table.cell(0, 1).text = "理解补充材料的具体经历"
        second.notes_slide.notes_text_frame.text = "访谈主题：一次补件经历"
        deck.save(stream)
    else:
        pdf = PdfWriter()
        page = pdf.add_blank_page(width=600, height=800)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): pdf._add_object(font)})}
        )
        content = DecodedStreamObject()
        content.set_data(b"BT /F1 12 Tf 50 750 Td (Research objective: understand document handoffs) Tj ET")
        page[NameObject("/Contents")] = pdf._add_object(content)
        pdf.write(stream)
    return stream.getvalue()


def upload(client, kind="docx", path="/api/admin/studies/import", content=None, extra=None):
    return client.post(
        path,
        content=content if content is not None else document_bytes(kind),
        headers={
            **headers(),
            "X-File-Name": quote(f"调研方案.{kind}"),
            "X-Import-Consent": "accepted",
            "X-Import-Version": client.get("/api/config").json()["consent_version"],
            **(extra or {}),
        },
    )


@pytest.mark.parametrize("kind", ["pdf", "docx", "xlsx", "pptx"])
def test_upload_reads_document_and_only_publishes_after_review(admin, app, kind):
    response = upload(admin, kind)
    assert response.status_code == 200, response.text
    sid, jid = response.json()["study_id"], response.json()["job_id"]
    study = admin.get(f"/api/admin/studies/{sid}").json()
    assert study["current_version_id"] is None
    assert admin.post(f"/api/admin/studies/{sid}/invites", json={}, headers=headers()).status_code == 409
    with app.state.db.read() as db:
        text = db.get(Job, jid).payload["document"]["text"]
        assert ("document handoffs" if kind == "pdf" else "理解补充材料的具体经历") in text
        if kind == "docx":
            assert text.index("材料流转体验") < text.index("理解补充材料") < text.index("一次补件经历")
        if kind == "xlsx":
            assert "访谈提纲" in text and "当时如何处理" in text
        if kind == "pptx":
            assert "幻灯片 2" in text and "一次补件经历" in text
    asyncio.run(Worker(app.state.db, app.state.settings).run_once())
    study = admin.get(f"/api/admin/studies/{sid}").json()
    assert study["import_job"]["status"] == "succeeded"
    assert study["import_job"]["result"]["mode"] == "mock"
    assert study["version"]["topics"]
    assert study["current_version_id"] is None
    result = admin.post(f"/api/admin/studies/{sid}/versions", json=study["version"], headers=headers())
    assert result.status_code == 200, result.text
    assert result.json()["version_number"] == 1
    assert admin.post(f"/api/admin/studies/{sid}/invites", json={}, headers=headers()).status_code == 200


def test_import_rejects_unreadable_or_unapproved_input_without_creating_study(admin, app):
    for kind, content, extra in [
        ("doc", b"old binary", {}),
        ("docx", b"not a document", {}),
        ("docx", document_bytes("docx"), {"X-Import-Consent": ""}),
    ]:
        assert upload(admin, kind, content=content, extra=extra).status_code in {400, 403, 415, 422}
    with app.state.db.read() as db:
        assert list(db.scalars(select(Study))) == []
        assert list(db.scalars(select(Job))) == []


def test_import_retains_source_across_retry_and_study_delete(admin, app):
    result = upload(admin).json()
    sid, jid = result["study_id"], result["job_id"]

    class Provider:
        async def text(self, role, messages, **kwargs):
            payload = json.loads(messages[-1]["content"])
            assert "理解补充材料的具体经历" in payload["document"]["text"]
            return TextResult(
                json.dumps({"study": STUDY, "warnings": ["核对目标时长"]}, ensure_ascii=False),
                Usage(source="mock"),
            )

    worker = Worker(app.state.db, app.state.settings, provider=Provider())
    asyncio.run(worker.run_once())
    study = admin.get(f"/api/admin/studies/{sid}").json()
    assert study["version"]["objective"] == "了解真实流程经历"
    assert study["import_source"]["sha256"] and "理解补充材料" in study["import_source"]["text"]
    assert admin.delete(f"/api/admin/studies/{sid}", headers=headers()).status_code == 200
    with app.state.db.read() as db:
        assert db.get(Job, jid) is None
        assert list(db.scalars(select(StudyVersion))) == []


def test_existing_study_import_does_not_modify_published_version(admin, app):
    original = new_study(admin)
    response = upload(admin, path=f"/api/admin/studies/{original['id']}/import")
    assert response.status_code == 200, response.text
    asyncio.run(Worker(app.state.db, app.state.settings).run_once())
    current = admin.get(f"/api/admin/studies/{original['id']}").json()
    assert current["version"] == original["version"]
    assert current["current_version_id"] == original["current_version_id"]
    assert current["import_job"]["result"]["study"]["topics"]


def test_duplicate_upload_is_idempotent_and_unauthenticated_upload_is_rejected(admin, app):
    from fastapi.testclient import TestClient

    request_headers = headers()
    content = document_bytes("docx")
    first = upload(admin, extra=request_headers, content=content)
    second = upload(admin, extra=request_headers, content=content)
    assert first.json() == second.json()
    assert upload(TestClient(app)).status_code == 401
    with app.state.db.read() as db:
        assert len(list(db.scalars(select(Study)))) == 1
        assert len(list(db.scalars(select(Job)))) == 1


def test_failed_model_result_can_be_retried_without_reupload_or_changed_source(admin, app):
    result = upload(admin).json()
    sid, jid = result["study_id"], result["job_id"]

    class Provider:
        calls = 0
        documents = []

        async def text(self, role, messages, **kwargs):
            self.calls += 1
            self.documents.append(json.loads(messages[-1]["content"])["document"])
            return TextResult(
                "invalid json" if self.calls == 1 else json.dumps({"study": STUDY}), Usage(source="mock")
            )

    provider = Provider()
    worker = Worker(app.state.db, app.state.settings, provider=provider)
    asyncio.run(worker.run_once())
    assert admin.get(f"/api/admin/jobs/{jid}").json()["error_code"] == "INVALID_STUDY_IMPORT"
    assert (
        admin.post(f"/api/admin/studies/{sid}/imports/{jid}/retry", json={}, headers=headers()).status_code
        == 200
    )
    asyncio.run(worker.run_once())
    assert admin.get(f"/api/admin/jobs/{jid}").json()["status"] == "succeeded"
    assert provider.calls == 2
    assert provider.documents[0] == provider.documents[1]


def test_unknown_charge_retry_requires_explicit_confirmation(admin, app):
    from app.providers import ProviderError

    result = upload(admin).json()
    sid, jid = result["study_id"], result["job_id"]

    class Provider:
        async def text(self, *args, **kwargs):
            raise ProviderError("EXTERNAL_STATUS_UNKNOWN", "结果未知", True)

    asyncio.run(Worker(app.state.db, app.state.settings, provider=Provider()).run_once())
    assert admin.get(f"/api/admin/jobs/{jid}").json()["status"] == "external_status_unknown"
    retry = f"/api/admin/studies/{sid}/imports/{jid}/retry"
    assert admin.post(retry, json={}, headers=headers()).status_code == 409
    assert upload(admin, path=f"/api/admin/studies/{sid}/import").status_code == 409
    assert admin.post(retry, json={"accept_possible_charge": True}, headers=headers()).status_code == 200
    asyncio.run(Worker(app.state.db, app.state.settings).run_once())
    assert admin.get(f"/api/admin/jobs/{jid}").json()["status"] == "succeeded"


def test_deleting_study_during_import_cannot_recreate_draft(admin, app):
    result = upload(admin).json()
    sid, jid = result["study_id"], result["job_id"]

    class Provider:
        async def text(self, *args, **kwargs):
            from app.storage import Storage

            Storage(app.state.db, app.state.settings).delete_study(sid)
            return TextResult(json.dumps({"study": STUDY}), Usage(source="mock"))

    asyncio.run(Worker(app.state.db, app.state.settings, provider=Provider()).run_once())
    with app.state.db.read() as db:
        assert db.get(Study, sid) is None
        assert db.get(Job, jid) is None


def test_backup_restores_import_source_and_filters_deleted_studies(admin, app, tmp_path):
    import sqlite3
    from app.storage import Storage

    kept, removed = upload(admin).json(), upload(admin, "pptx").json()
    storage = Storage(app.state.db, app.state.settings)
    backup = storage.backup()
    storage.delete_study(removed["study_id"])
    target = tmp_path / "restored-import"
    storage.restore(backup["backup_path"], target)
    with sqlite3.connect(target / "probeflow.sqlite3") as db:
        rows = db.execute("SELECT id, payload FROM jobs WHERE kind = ?", ("study_import",)).fetchall()
    assert len(rows) == 1 and rows[0][0] == kept["job_id"]
    assert "理解补充材料的具体经历" in json.loads(rows[0][1])["document"]["text"]


@pytest.mark.parametrize("protected", [False, True])
def test_empty_and_encrypted_pdf_do_not_create_placeholder_research(admin, app, protected):
    writer, output = PdfWriter(), io.BytesIO()
    writer.add_blank_page(width=300, height=300)
    if protected:
        writer.encrypt("fictional-password")
    writer.write(output)
    response = upload(admin, "pdf", content=output.getvalue())
    assert response.status_code == 422
    assert response.json()["code"] == ("DOCUMENT_PROTECTED" if protected else "DOCUMENT_NO_TEXT")
    with app.state.db.read() as db:
        assert list(db.scalars(select(Study))) == []


def test_oversized_extracted_text_is_rejected_instead_of_silently_cut(admin, app):
    doc, output = Document(), io.BytesIO()
    doc.add_paragraph("虚构方案" * 11000)
    doc.save(output)
    response = upload(admin, content=output.getvalue())
    assert response.status_code == 413
    assert response.json()["code"] == "DOCUMENT_TOO_LONG"
    assert admin.get("/api/admin/studies").json() == []


def test_nested_word_table_keeps_questions_and_exclusions(admin, app):
    document, stream = Document(), io.BytesIO()
    document.add_paragraph("虚构调研方案：材料流转体验")
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "访谈详细方案"
    nested = cell.add_table(rows=2, cols=1)
    nested.cell(0, 0).text = "主题一：补件经历"
    nested.cell(1, 0).text = "绝对不可询问客户身份及账号"
    document.save(stream)
    response = upload(admin, content=stream.getvalue())
    assert response.status_code == 200
    source = admin.get(f"/api/admin/studies/{response.json()['study_id']}").json()["import_source"]["text"]
    assert "主题一：补件经历" in source
    assert "绝对不可询问客户身份及账号" in source


def test_changed_processing_configuration_cannot_send_a_previously_consented_import(admin, app):
    result = upload(admin).json()
    app.state.settings.mode = "live"

    class Provider:
        calls = 0

        async def text(self, *args, **kwargs):
            self.calls += 1
            return TextResult(json.dumps({"study": STUDY}), Usage(source="mock"))

    provider = Provider()
    asyncio.run(Worker(app.state.db, app.state.settings, provider=provider).run_once())
    assert provider.calls == 0
    job = admin.get(f"/api/admin/jobs/{result['job_id']}").json()
    assert job["error_code"] == "IMPORT_CONSENT_CHANGED"


def test_unknown_import_can_be_replaced_after_configuration_change_with_fresh_charge_consent(admin, app):
    original = upload(admin).json()
    with app.state.db.transaction() as db:
        db.get(Job, original["job_id"]).status = "external_status_unknown"
    app.state.settings.interview_model = "different-model"
    path = f"/api/admin/studies/{original['study_id']}/import"
    assert upload(admin, path=path).status_code == 409
    replacement = upload(admin, path=path, extra={"X-Accept-Possible-Charge": "true"})
    assert replacement.status_code == 200, replacement.text
    assert replacement.json()["job_id"] != original["job_id"]
    with app.state.db.read() as db:
        assert db.get(Job, original["job_id"]).status == "external_status_unknown"
        assert (
            db.get(Job, replacement.json()["job_id"]).payload["processing_snapshot"]["provider"]["model"]
            == "different-model"
        )
