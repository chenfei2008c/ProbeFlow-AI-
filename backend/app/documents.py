"""Bounded, local text extraction. Never execute Office macros or fetch external links."""

import hashlib
import io
import json
import re
import subprocess
import sys
import zipfile
from pathlib import Path

from app.errors import AppError

MAX_BYTES = 10 * 1024 * 1024
MAX_CHARS = 40000


def extract_document(filename: str, content: bytes) -> dict:
    filename = Path(filename.replace("\\", "/")).name
    if not filename or len(filename) > 240 or any(ord(c) < 32 for c in filename):
        raise AppError("DOCUMENT_NAME_INVALID", "文件名无效，请重命名后上传。")
    kind = Path(filename).suffix.lower().lstrip(".")
    if kind not in {"pdf", "docx", "xlsx", "pptx"}:
        raise AppError(
            "DOCUMENT_FORMAT", "支持 PDF、DOCX、XLSX、PPTX；旧版 DOC、XLS、PPT 请先另存为对应新格式。", 415
        )
    if not content or len(content) > MAX_BYTES:
        raise AppError("DOCUMENT_SIZE", "请上传非空且不超过 10 MB 的文件。", 413)
    parts, warnings = [], []
    total = 0

    def add(text):
        nonlocal total
        text = str(text).strip()
        if not text:
            return
        total += len(text) + 1
        if total > MAX_CHARS:
            raise AppError(
                "DOCUMENT_TOO_LONG", "提取内容超过 40,000 字符，请拆分方案后上传；未截断内容。", 413
            )
        parts.append(text)

    try:
        if kind != "pdf":
            with zipfile.ZipFile(io.BytesIO(content)) as package:
                entries = package.infolist()
                if len(entries) > 3000 or sum(e.file_size for e in entries) > 40 * 1024 * 1024:
                    raise AppError("DOCUMENT_SIZE", "文件解压后过大，请精简后上传。", 413)
                if any(e.flag_bits & 1 or "vbaProject" in e.filename for e in entries):
                    raise AppError("DOCUMENT_PROTECTED", "请上传未加密、不含宏的文档。", 422)
        if kind == "docx":
            from docx import Document
            from docx.table import Table

            def body_text(container):
                for item in container.iter_inner_content():
                    if isinstance(item, Table):
                        for row in item.rows:
                            cells, seen = [], set()
                            for cell in row.cells:
                                if cell._tc not in seen:
                                    seen.add(cell._tc)
                                    cells.append(" / ".join(body_text(cell)))
                            yield " | ".join(cells)
                    else:
                        yield item.text

            for text in body_text(Document(io.BytesIO(content))):
                add(text)
        elif kind == "xlsx":
            from openpyxl import load_workbook

            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                if len(book.worksheets) > 30:
                    raise AppError("DOCUMENT_SIZE", "工作表超过 30 张，请拆分后上传。", 413)
                cells_read = 0
                for sheet in book.worksheets:
                    add(f"【工作表：{sheet.title}】")
                    if (sheet.max_row or 0) > 10000 or (sheet.max_column or 0) > 200:
                        raise AppError("DOCUMENT_SIZE", "工作表范围过大，请仅保留调研方案内容后上传。", 413)
                    sheet.reset_dimensions()
                    for row in sheet.iter_rows():
                        cells_read += len(row)
                        if cells_read > 50000:
                            raise AppError("DOCUMENT_SIZE", "单元格超过处理上限，请精简后上传。", 413)
                        values = []
                        for cell in row:
                            if cell.value is not None:
                                if cell.data_type == "f":
                                    warnings.append(
                                        "Excel 公式保留为公式文本，未执行计算；请核对涉及计算的内容。"
                                    )
                                values.append(f"{cell.coordinate}: {cell.value}")
                        add(" | ".join(values))
            finally:
                book.close()
        elif kind == "pptx":
            from pptx import Presentation
            from pptx.enum.shapes import MSO_SHAPE_TYPE

            deck = Presentation(io.BytesIO(content))
            if len(deck.slides) > 100:
                raise AppError("DOCUMENT_SIZE", "幻灯片超过 100 页，请拆分后上传。", 413)

            def read_shapes(shapes):
                for shape in shapes:
                    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                        read_shapes(shape.shapes)
                    if shape.has_text_frame:
                        add(shape.text)
                    if shape.has_table:
                        for row in shape.table.rows:
                            add(" | ".join(cell.text for cell in row.cells))
                    if shape.has_chart:
                        warnings.append("PPT 图表及图片中的文字未提取，请核对提取原文。")

            for number, slide in enumerate(deck.slides, 1):
                add(f"【幻灯片 {number}】")
                read_shapes(slide.shapes)
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    add(slide.notes_slide.notes_text_frame.text)
        else:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted:
                raise AppError("DOCUMENT_PROTECTED", "PDF 已加密，请先解除密码后上传。", 422)
            if len(reader.pages) > 100:
                raise AppError("DOCUMENT_SIZE", "PDF 超过 100 页，请拆分后上传。", 413)
            for number, page in enumerate(reader.pages, 1):
                stream = page.get_contents()
                if stream is not None and len(stream.get_data()) > 5 * 1024 * 1024:
                    raise AppError("DOCUMENT_SIZE", "PDF 单页内容过于复杂，请精简或另存为 Word 后上传。", 413)
                text = page.extract_text() or ""
                if not text.strip():
                    warnings.append(f"PDF 第 {number} 页没有可提取文字，可能是空白页或图片。")
                else:
                    add(f"【PDF 第 {number} 页】\n{text}")
    except AppError:
        raise
    except Exception as exc:
        raise AppError(
            "DOCUMENT_UNREADABLE", "无法读取文件，请确认格式、文件完整性及是否已解除密码。", 422
        ) from exc
    text = "\n".join(parts)
    substantive = re.sub(r"【[^】]*】", "", text).strip()
    if len(substantive) < 10:
        raise AppError(
            "DOCUMENT_NO_TEXT",
            "没有足够的可读取文字。扫描件或图片请先转为可复制文字的 PDF／Office 文档。",
            422,
        )
    return {
        "filename": filename,
        "format": kind,
        "sha256": hashlib.sha256(content).hexdigest(),
        "text": text,
        "warnings": list(dict.fromkeys(warnings)),
        "characters": len(text),
    }


def parse_document(filename: str, content: bytes) -> dict:
    """Isolate parsers from the application process and enforce a wall-clock deadline."""
    try:
        process = subprocess.run(
            [sys.executable, "-m", "app.documents", filename],
            input=content,
            cwd=Path(__file__).resolve().parents[1],
            stdout=subprocess.PIPE,
            stderr=subprocess.DEVNULL,
            timeout=25,
            check=False,
        )
        result = json.loads(process.stdout) if process.returncode == 0 else None
    except (subprocess.TimeoutExpired, ValueError):
        result = None
    if not result:
        raise AppError("DOCUMENT_PARSE_LIMIT", "文件解析超时或资源超限，请精简文件后重试。", 422)
    if "error" in result:
        raise AppError(**result["error"])
    return result


if __name__ == "__main__":
    # On supported systems, bound parser CPU and address space as well as parent timeout.
    try:
        import resource

        resource.setrlimit(resource.RLIMIT_CPU, (15, 15))
        if sys.platform.startswith("linux"):
            resource.setrlimit(resource.RLIMIT_AS, (768 * 1024 * 1024,) * 2)
    except (ImportError, OSError, ValueError):
        pass
    try:
        response = extract_document(sys.argv[1], sys.stdin.buffer.read(MAX_BYTES + 1))
    except AppError as exc:
        response = {"error": {"code": exc.code, "message": exc.message, "status": exc.status}}
    print(json.dumps(response, ensure_ascii=False))
